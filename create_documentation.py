from docx import Document
from docx.shared import Inches, Pt, Cm
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.style import WD_STYLE_TYPE

# Create document
doc = Document()

# Set up styles
style = doc.styles['Normal']
font = style.font
font.name = 'Times New Roman'
font.size = Pt(12)

# ==================== TITLE PAGE ====================
# Add some spacing at top
for _ in range(4):
    doc.add_paragraph()

# University name
uni = doc.add_paragraph()
uni.alignment = WD_ALIGN_PARAGRAPH.CENTER
run = uni.add_run("POLITEHNICA UNIVERSITY OF TIMIȘOARA")
run.bold = True
run.font.size = Pt(16)
run.font.name = 'Times New Roman'

# Faculty
faculty = doc.add_paragraph()
faculty.alignment = WD_ALIGN_PARAGRAPH.CENTER
run = faculty.add_run("Faculty of Automation and Computer Science")
run.font.size = Pt(14)
run.font.name = 'Times New Roman'

# Add spacing
for _ in range(6):
    doc.add_paragraph()

# Project title
title = doc.add_paragraph()
title.alignment = WD_ALIGN_PARAGRAPH.CENTER
run = title.add_run("PROJECT DOCUMENTATION")
run.bold = True
run.font.size = Pt(18)
run.font.name = 'Times New Roman'

# App name
app_name = doc.add_paragraph()
app_name.alignment = WD_ALIGN_PARAGRAPH.CENTER
run = app_name.add_run("News App")
run.bold = True
run.font.size = Pt(24)
run.font.name = 'Times New Roman'

# Subtitle
subtitle = doc.add_paragraph()
subtitle.alignment = WD_ALIGN_PARAGRAPH.CENTER
run = subtitle.add_run("Android Application for Displaying News")
run.italic = True
run.font.size = Pt(14)
run.font.name = 'Times New Roman'

# Add spacing
for _ in range(8):
    doc.add_paragraph()

# Author info
author = doc.add_paragraph()
author.alignment = WD_ALIGN_PARAGRAPH.RIGHT
run = author.add_run("Student: Lazar Razvan")
run.font.size = Pt(12)
run.font.name = 'Times New Roman'

# Date
date = doc.add_paragraph()
date.alignment = WD_ALIGN_PARAGRAPH.RIGHT
run = date.add_run("January 2026")
run.font.size = Pt(12)
run.font.name = 'Times New Roman'

# Page break
doc.add_page_break()

# ==================== TABLE OF CONTENTS ====================
toc_title = doc.add_heading('Table of Contents', level=1)

toc_items = [
    ("1. Introduction", "3"),
    ("2. Application Description", "3"),
    ("3. Technologies Used", "4"),
    ("4. Project Structure", "5"),
    ("5. Application Functionality", "6"),
    ("6. User Interface", "7"),
    ("7. News API", "8"),
    ("8. Conclusions", "9"),
]

for item, page in toc_items:
    p = doc.add_paragraph()
    p.add_run(item)
    p.add_run("\t" * 8 + page)

doc.add_page_break()

# ==================== CONTENT ====================

# 1. INTRODUCTION
doc.add_heading('1. Introduction', level=1)

doc.add_paragraph(
    "This documentation presents the News App project, a modern Android application "
    "developed for displaying news in real-time. The application was created as a learning "
    "project to explore modern Android development technologies, specifically Kotlin and "
    "Jetpack Compose."
)

doc.add_paragraph(
    "The main purpose of this project was to understand how Android application development "
    "works using the latest tools and libraries available. The application demonstrates "
    "important concepts such as MVVM architecture, reactive programming with Kotlin Flow, "
    "and integration with web APIs."
)

# 2. APPLICATION DESCRIPTION
doc.add_heading('2. Application Description', level=1)

doc.add_paragraph(
    "News App is an application that connects to the internet to download and display "
    "the latest news from around the world. Users can browse through the list of news "
    "articles and tap on any article to read the complete story in their web browser."
)

doc.add_heading('Main Features', level=2)

features = [
    "Real-time Display - The application downloads the latest news headlines immediately when opened",
    "Modern Design - The user interface uses Material Design 3, Google's latest design system",
    "Attractive Visual Cards - Each article is displayed in a card with image, title, description, and source",
    "Read Full Articles - Tapping a card opens the complete article in the browser",
    "Automatic Themes - The application supports light mode and dark mode, adapting to phone settings",
    "Dynamic Colors - On Android 12+, the application colors automatically match your wallpaper"
]

for feature in features:
    p = doc.add_paragraph(feature, style='List Bullet')

# 3. TECHNOLOGIES USED
doc.add_heading('3. Technologies Used', level=1)

doc.add_paragraph(
    "For developing this application, I used a modern set of technologies and libraries "
    "that represent the current standard in Android development:"
)

doc.add_heading('Kotlin', level=2)
doc.add_paragraph(
    "Kotlin is the main programming language used for Android development. It is a modern, "
    "concise, and safe language officially recommended by Google for Android applications. "
    "Kotlin offers many advantages over Java, including simpler syntax, null safety, and "
    "coroutines for asynchronous programming."
)

doc.add_heading('Jetpack Compose', level=2)
doc.add_paragraph(
    "Jetpack Compose is Google's modern toolkit for building native Android user interfaces. "
    "Instead of using XML files for layouts, the entire interface is written in Kotlin using "
    "composable functions. This makes the code easier to understand and maintain."
)

doc.add_heading('Material Design 3', level=2)
doc.add_paragraph(
    "Material Design 3 (also known as Material You) is the latest design system from Google. "
    "It provides beautiful and consistent UI components, support for dynamic themes, and a "
    "modern visual experience."
)

doc.add_heading('Retrofit', level=2)
doc.add_paragraph(
    "Retrofit is a popular library for communicating with web APIs. It makes it very easy "
    "to download data from the internet and convert it to Kotlin objects. In this application, "
    "Retrofit is used to get news from NewsAPI."
)

doc.add_heading('Coil', level=2)
doc.add_paragraph(
    "Coil (Coroutine Image Loader) is a modern library for loading images in Android. "
    "It is optimized for Kotlin and Jetpack Compose, being fast and memory-efficient. "
    "We use it to display images from news articles."
)

doc.add_heading('Kotlin Coroutines and Flow', level=2)
doc.add_paragraph(
    "Kotlin Coroutines allow executing long-running operations (such as downloading data "
    "from the internet) without blocking the user interface. Flow is a reactive feature "
    "that allows automatic UI updates when data changes."
)

doc.add_page_break()

# 4. PROJECT STRUCTURE
doc.add_heading('4. Project Structure', level=1)

doc.add_paragraph(
    "The project is organized in a clear and maintainable structure, following the "
    "MVVM (Model-View-ViewModel) architecture principles:"
)

doc.add_heading('Main File', level=2)
doc.add_paragraph(
    "MainActivity.kt - This is the entry point of the application. This is where the "
    "theme and main screen of the application are initialized."
)

doc.add_heading('Data Folder', level=2)
doc.add_paragraph(
    "This folder contains everything related to obtaining data from external sources:"
)
data_files = [
    "NewsApi.kt - Defines the interface for communicating with the news API and the data structure (Article, Source, NewsResponse)",
    "NewsRepository.kt - Manages the internet connection and news downloading, using Retrofit for communication"
]
for item in data_files:
    doc.add_paragraph(item, style='List Bullet')

doc.add_heading('UI Folder', level=2)
doc.add_paragraph(
    "This folder contains everything the user sees and interacts with:"
)
ui_files = [
    "NewsScreen.kt - The main screen with the news list, includes the ArticleItem component for displaying each article",
    "NewsViewModel.kt - Manages the business logic and data state for the screen"
]
for item in ui_files:
    doc.add_paragraph(item, style='List Bullet')

doc.add_heading('Theme Folder', level=2)
doc.add_paragraph(
    "Contains the visual styling of the application:"
)
theme_files = [
    "Color.kt - The color palette used in the application",
    "Theme.kt - Configuration for light mode, dark mode, and dynamic color themes",
    "Type.kt - Typography styles for text"
]
for item in theme_files:
    doc.add_paragraph(item, style='List Bullet')

# 5. APPLICATION FUNCTIONALITY
doc.add_heading('5. Application Functionality', level=1)

doc.add_heading('Application Startup', level=2)
doc.add_paragraph(
    "When the user opens the application, it automatically starts the process of downloading "
    "news. A loading indicator (spinner) appears in the center of the screen to show that "
    "data is being downloaded."
)

doc.add_heading('Downloading News', level=2)
doc.add_paragraph(
    "The application connects to NewsAPI.org using the configured API key. It sends a request "
    "to get the top headlines from the United States. The response includes a list of articles, "
    "each containing: title, description, image URL, full article URL, and source name."
)

doc.add_heading('Displaying the List', level=2)
doc.add_paragraph(
    "After the news is downloaded, the loading indicator disappears and the list of articles "
    "appears. The list is implemented with LazyColumn, an efficient component that loads only "
    "the elements visible on the screen, saving memory."
)

doc.add_heading('User Interaction', level=2)
doc.add_paragraph(
    "The user can scroll through the list of news. When they tap on an article card, "
    "the application opens the default web browser and navigates to the original article URL, "
    "where the user can read the entire story."
)

doc.add_page_break()

# 6. USER INTERFACE
doc.add_heading('6. User Interface', level=1)

doc.add_heading('Top App Bar', level=2)
doc.add_paragraph(
    "At the top of the screen is a bar with the title \"Top Headlines\". This uses the "
    "Material 3 style and adapts to the active theme (light or dark)."
)

doc.add_heading('Article Card', level=2)
doc.add_paragraph(
    "Each news article is displayed in a card that contains:"
)
card_elements = [
    "Image - The main photo of the article, displayed at the top of the card with a height of 200dp",
    "Title - The article headline, displayed in bold font to stand out",
    "Description - A short summary of the article, limited to a maximum of 3 lines to maintain visual consistency",
    "Source - The name of the publication that published the article, displayed in the theme's primary color"
]
for item in card_elements:
    doc.add_paragraph(item, style='List Bullet')

doc.add_heading('Loading Indicator', level=2)
doc.add_paragraph(
    "While the application downloads data from the internet, a CircularProgressIndicator "
    "appears in the center of the screen. It automatically disappears when the news is "
    "ready to be displayed."
)

# 7. NEWS API
doc.add_heading('7. News API', level=1)

doc.add_paragraph(
    "The application uses NewsAPI.org, a free service for developers that provides "
    "access to news articles from thousands of sources around the world."
)

doc.add_heading('Endpoint Used', level=2)
doc.add_paragraph(
    "The application accesses the /v2/top-headlines endpoint to get the most important "
    "headlines. The request includes parameters for country (us) and the API key for "
    "authentication."
)

doc.add_heading('Response Structure', level=2)
doc.add_paragraph("The API response contains:")
response_items = [
    "status - Indicates whether the request was processed successfully",
    "totalResults - The total number of articles available",
    "articles - The list of articles, each with title, description, URL, image, and source"
]
for item in response_items:
    doc.add_paragraph(item, style='List Bullet')

doc.add_heading('Limitations', level=2)
doc.add_paragraph(
    "The free plan of NewsAPI has some limitations: maximum 100 requests per day, "
    "only works during development (not for published apps), and news may be delayed "
    "by a few hours. For a production application, upgrading to a paid plan would be "
    "necessary."
)

# 8. CONCLUSIONS
doc.add_heading('8. Conclusions', level=1)

doc.add_paragraph(
    "Developing this application was a valuable learning experience. I successfully "
    "implemented a functional Android application using the most modern technologies "
    "available in the Android ecosystem."
)

doc.add_paragraph(
    "Through this project I learned:"
)

lessons = [
    "How to use Kotlin for Android development",
    "How to build user interfaces with Jetpack Compose",
    "How to connect the application to a web API using Retrofit",
    "How to manage application state using ViewModel and Flow",
    "How to load and display images efficiently with Coil",
    "How to implement themes and styles with Material Design 3"
]
for lesson in lessons:
    doc.add_paragraph(lesson, style='List Bullet')

doc.add_paragraph(
    "The application can be improved in the future with features such as: news search, "
    "category filtering, saving favorite articles, offline mode, and support for "
    "multiple countries."
)

doc.add_paragraph()
final = doc.add_paragraph()
final.alignment = WD_ALIGN_PARAGRAPH.CENTER
run = final.add_run("— End of Documentation —")
run.italic = True

# Save document
doc.save('NewsApp_Documentation.docx')
print("Documentation created successfully: NewsApp_Documentation.docx")
