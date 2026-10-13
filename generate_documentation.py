"""
NewsApp Complete Documentation Generator
Creates a comprehensive Word document with full project documentation
"""

from docx import Document
from docx.shared import Inches, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.style import WD_STYLE_TYPE
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.oxml.ns import qn
from docx.oxml import OxmlElement

# Create document
doc = Document()

# Helper functions
def set_cell_shading(cell, color):
    """Set cell background color"""
    shading = OxmlElement('w:shd')
    shading.set(qn('w:fill'), color)
    cell._tc.get_or_add_tcPr().append(shading)

def add_heading_styled(text, level=1):
    """Add a styled heading"""
    heading = doc.add_heading(text, level)
    if level == 1:
        heading.runs[0].font.color.rgb = RGBColor(79, 70, 229)
    return heading

def add_code_block(code, language="kotlin"):
    """Add a formatted code block"""
    # Add a paragraph with code styling
    para = doc.add_paragraph()
    run = para.add_run(code)
    run.font.name = 'Consolas'
    run.font.size = Pt(9)
    para.paragraph_format.left_indent = Inches(0.25)
    para.paragraph_format.space_before = Pt(6)
    para.paragraph_format.space_after = Pt(6)
    # Add light gray background effect via paragraph shading
    shading = OxmlElement('w:shd')
    shading.set(qn('w:fill'), 'F1F5F9')
    para._p.get_or_add_pPr().append(shading)
    return para

# ============================================================
# TITLE PAGE
# ============================================================
title = doc.add_heading('News App', 0)
title.alignment = WD_ALIGN_PARAGRAPH.CENTER
title.runs[0].font.size = Pt(48)
title.runs[0].font.color.rgb = RGBColor(79, 70, 229)

subtitle = doc.add_paragraph('Complete Technical Documentation')
subtitle.alignment = WD_ALIGN_PARAGRAPH.CENTER
subtitle.runs[0].font.size = Pt(24)
subtitle.runs[0].font.color.rgb = RGBColor(100, 116, 139)

doc.add_paragraph()
doc.add_paragraph()

info_para = doc.add_paragraph()
info_para.alignment = WD_ALIGN_PARAGRAPH.CENTER
info_para.add_run('A Modern Android News Application\n').bold = True
info_para.add_run('Built with Kotlin & Jetpack Compose\n\n')
info_para.add_run('Version 1.0')

doc.add_page_break()

# ============================================================
# TABLE OF CONTENTS
# ============================================================
add_heading_styled('Table of Contents', 1)

toc_items = [
    ('1. Introduction', '3'),
    ('2. Project Overview', '3'),
    ('3. Technologies Used', '4'),
    ('4. Project Structure', '5'),
    ('5. Architecture', '6'),
    ('6. Core Components', '7'),
    ('   6.1 MainActivity', '7'),
    ('   6.2 Data Layer', '8'),
    ('   6.3 UI Layer', '10'),
    ('   6.4 Theme', '12'),
    ('7. API Integration', '13'),
    ('8. User Interface', '14'),
    ('9. Configuration & Setup', '15'),
    ('10. Future Improvements', '16'),
    ('11. Appendix', '17'),
]

toc_table = doc.add_table(rows=len(toc_items), cols=2)
for i, (item, page) in enumerate(toc_items):
    toc_table.rows[i].cells[0].text = item
    toc_table.rows[i].cells[1].text = page
    toc_table.rows[i].cells[1].paragraphs[0].alignment = WD_ALIGN_PARAGRAPH.RIGHT

doc.add_page_break()

# ============================================================
# 1. INTRODUCTION
# ============================================================
add_heading_styled('1. Introduction', 1)

doc.add_paragraph(
    'News App is a modern Android application designed to fetch and display top headlines from around '
    'the world. The application provides users with real-time access to news articles from various sources, '
    'presented in an elegant and user-friendly interface.'
)

doc.add_paragraph(
    'This documentation provides a comprehensive overview of the application architecture, code structure, '
    'technologies used, and implementation details. It is intended for developers who want to understand, '
    'maintain, or extend the application.'
)

add_heading_styled('Purpose', 2)
doc.add_paragraph(
    'The primary purpose of this application is to demonstrate modern Android development practices using:'
)
bullets = [
    'Kotlin as the primary programming language',
    'Jetpack Compose for declarative UI development',
    'MVVM (Model-View-ViewModel) architecture pattern',
    'Coroutines for asynchronous programming',
    'Retrofit for network operations',
]
for bullet in bullets:
    p = doc.add_paragraph(bullet, style='List Bullet')

# ============================================================
# 2. PROJECT OVERVIEW
# ============================================================
add_heading_styled('2. Project Overview', 1)

add_heading_styled('Application Details', 2)

# App details table
details_table = doc.add_table(rows=7, cols=2)
details_table.style = 'Table Grid'
details = [
    ('Application Name', 'News App'),
    ('Package Name', 'com.example.newsapp'),
    ('Minimum SDK', '24 (Android 7.0)'),
    ('Target SDK', '34 (Android 14)'),
    ('Version', '1.0'),
    ('Language', 'Kotlin'),
    ('UI Framework', 'Jetpack Compose'),
]
for i, (key, value) in enumerate(details):
    details_table.rows[i].cells[0].text = key
    details_table.rows[i].cells[1].text = value
    set_cell_shading(details_table.rows[i].cells[0], 'F1F5F9')

doc.add_paragraph()

add_heading_styled('Key Features', 2)
features = [
    ('Live News Feed', 'Connects to the internet and downloads the latest top headlines in real-time from NewsAPI.org'),
    ('Beautiful Card UI', 'Each news article is displayed in a visually appealing card format with the article image, title, description, and source'),
    ('Read Full Articles', 'Users can tap on any news card to open the complete article in their default web browser'),
    ('Modern Design', 'Uses Google\'s latest Material 3 design system with support for both light and dark modes'),
    ('Dynamic Colors', 'On Android 12 and above, the app colors automatically match the phone\'s wallpaper'),
]

for title, desc in features:
    p = doc.add_paragraph()
    p.add_run(f'• {title}: ').bold = True
    p.add_run(desc)

doc.add_page_break()

# ============================================================
# 3. TECHNOLOGIES USED
# ============================================================
add_heading_styled('3. Technologies Used', 1)

doc.add_paragraph(
    'The application leverages a modern technology stack that represents current best practices in Android development:'
)

tech_table = doc.add_table(rows=7, cols=3)
tech_table.style = 'Table Grid'

# Header row
headers = ['Technology', 'Version', 'Purpose']
for i, header in enumerate(headers):
    tech_table.rows[0].cells[i].text = header
    tech_table.rows[0].cells[i].paragraphs[0].runs[0].bold = True
    set_cell_shading(tech_table.rows[0].cells[i], '4F46E5')
    tech_table.rows[0].cells[i].paragraphs[0].runs[0].font.color.rgb = RGBColor(255, 255, 255)

technologies = [
    ('Kotlin', '1.9.x', 'Primary programming language for Android development'),
    ('Jetpack Compose', '1.5.x', 'Modern declarative UI toolkit for building native Android interfaces'),
    ('Material 3', '1.2.x', 'Google\'s latest design system providing UI components and theming'),
    ('Retrofit', '2.9.x', 'Type-safe HTTP client for making API requests'),
    ('Coil', '2.5.x', 'Image loading library optimized for Kotlin and Compose'),
    ('Coroutines', '1.7.x', 'Library for asynchronous programming and background tasks'),
]

for i, (tech, version, purpose) in enumerate(technologies):
    tech_table.rows[i+1].cells[0].text = tech
    tech_table.rows[i+1].cells[1].text = version
    tech_table.rows[i+1].cells[2].text = purpose

doc.add_paragraph()

add_heading_styled('Dependencies', 2)
doc.add_paragraph('The following dependencies are declared in the app\'s build.gradle.kts file:')

deps_code = """dependencies {
    // Android Core
    implementation(libs.androidx.core.ktx)
    implementation(libs.androidx.lifecycle.runtime.ktx)
    implementation(libs.androidx.lifecycle.viewmodel.ktx)
    implementation(libs.androidx.activity.compose)
    
    // Compose
    implementation(platform(libs.androidx.compose.bom))
    implementation(libs.androidx.ui)
    implementation(libs.androidx.ui.graphics)
    implementation(libs.androidx.ui.tooling.preview)
    implementation(libs.androidx.material3)
    
    // Networking
    implementation(libs.retrofit)
    implementation(libs.retrofit.gson)
    implementation(libs.coil.compose)
    
    // Navigation
    implementation(libs.androidx.navigation.compose)
}"""
add_code_block(deps_code)

doc.add_page_break()

# ============================================================
# 4. PROJECT STRUCTURE
# ============================================================
add_heading_styled('4. Project Structure', 1)

doc.add_paragraph(
    'The project follows a clean, modular structure that separates concerns and makes the codebase maintainable:'
)

structure = """com.example.newsapp/
├── MainActivity.kt          # Application entry point
├── data/                     # Data layer
│   ├── NewsApi.kt           # API interface and data classes
│   └── NewsRepository.kt    # Repository for data operations
└── ui/                       # UI layer
    ├── NewsScreen.kt        # Main screen composable
    ├── NewsViewModel.kt     # ViewModel for state management
    └── theme/               # Theming
        ├── Color.kt         # Color definitions
        ├── Theme.kt         # Theme configuration
        └── Type.kt          # Typography settings"""

add_code_block(structure, "text")

add_heading_styled('Directory Descriptions', 2)

dirs = [
    ('data/', 'Contains all data-related code including API definitions, data models, and the repository pattern implementation for fetching news from the network.'),
    ('ui/', 'Contains all user interface components built with Jetpack Compose, including screens, ViewModels, and composable functions.'),
    ('ui/theme/', 'Contains theming configuration including colors, typography, and the main theme setup for Material 3 with dynamic colors support.'),
]

for dir_name, desc in dirs:
    p = doc.add_paragraph()
    p.add_run(f'{dir_name} ').bold = True
    p.add_run(f'- {desc}')

doc.add_page_break()

# ============================================================
# 5. ARCHITECTURE
# ============================================================
add_heading_styled('5. Architecture', 1)

doc.add_paragraph(
    'The application follows the MVVM (Model-View-ViewModel) architecture pattern, which is the recommended '
    'architecture for Android applications by Google. This pattern separates the application into three main layers:'
)

add_heading_styled('Architecture Layers', 2)

layers = [
    ('View (UI Layer)', 'Composable functions that display data and handle user interactions. The UI observes the ViewModel\'s state and reacts to changes automatically.'),
    ('ViewModel', 'Holds and manages UI-related data. Survives configuration changes and exposes data as observable StateFlow streams.'),
    ('Model (Data Layer)', 'Handles data operations including network requests through the Repository and API interfaces. Abstracts data sources from the rest of the application.'),
]

for layer, desc in layers:
    p = doc.add_paragraph()
    p.add_run(f'{layer}: ').bold = True
    p.add_run(desc)

add_heading_styled('Data Flow', 2)

doc.add_paragraph('The data flows through the application in the following manner:')

flow_steps = [
    'User opens the app, triggering MainActivity',
    'MainActivity creates the NewsViewModel',
    'NewsViewModel initializes and calls fetchNews()',
    'NewsRepository makes an HTTP request to NewsAPI.org via Retrofit',
    'API returns JSON response which is parsed into Article objects',
    'NewsViewModel updates its StateFlow with the article list',
    'NewsScreen observes the StateFlow and recomposes the UI',
    'Articles are displayed in a LazyColumn as card items',
    'User taps a card, opening the article URL in the browser',
]

for i, step in enumerate(flow_steps, 1):
    doc.add_paragraph(f'{i}. {step}')

doc.add_page_break()

# ============================================================
# 6. CORE COMPONENTS
# ============================================================
add_heading_styled('6. Core Components', 1)

# 6.1 MainActivity
add_heading_styled('6.1 MainActivity', 2)

doc.add_paragraph(
    'MainActivity serves as the entry point of the application. It sets up the Compose UI and initializes '
    'the ViewModel using Kotlin property delegation.'
)

doc.add_paragraph('File: ').add_run('MainActivity.kt').bold = True

main_activity_code = """package com.example.newsapp

class MainActivity : ComponentActivity() {
    private val viewModel: NewsViewModel by viewModels()

    override fun onCreate(savedInstanceState: Bundle?) {
        super.onCreate(savedInstanceState)
        setContent {
            NewsAppTheme {
                Surface(
                    modifier = Modifier.fillMaxSize(),
                    color = MaterialTheme.colorScheme.background
                ) {
                    NewsScreen(viewModel = viewModel)
                }
            }
        }
    }
}"""
add_code_block(main_activity_code)

doc.add_paragraph('Key Points:')
points = [
    'Uses ComponentActivity for Compose support',
    'ViewModel is created using the viewModels() delegate for proper lifecycle handling',
    'NewsAppTheme wraps the content for consistent theming',
    'Surface provides the background color from the theme',
]
for point in points:
    doc.add_paragraph(point, style='List Bullet')

doc.add_page_break()

# 6.2 Data Layer
add_heading_styled('6.2 Data Layer', 2)

add_heading_styled('NewsApi.kt', 3)

doc.add_paragraph(
    'Defines the Retrofit interface for communicating with NewsAPI.org and the data classes for parsing '
    'the JSON response.'
)

news_api_code = """package com.example.newsapp.data

interface NewsApi {
    @GET("top-headlines")
    suspend fun getTopHeadlines(
        @Query("country") country: String = "us",
        @Query("apiKey") apiKey: String
    ): NewsResponse
}

data class NewsResponse(
    val status: String,
    val totalResults: Int,
    val articles: List<Article>
)

data class Article(
    val source: Source,
    val author: String?,
    val title: String,
    val description: String?,
    val url: String,
    val urlToImage: String?,
    val publishedAt: String,
    val content: String?
)

data class Source(
    val id: String?,
    val name: String
)"""
add_code_block(news_api_code)

doc.add_paragraph('Key Components:')
api_points = [
    'NewsApi interface: Defines the API endpoints using Retrofit annotations',
    'suspend function: Enables coroutine-based asynchronous calls',
    'NewsResponse: Wrapper class for the API response',
    'Article: Data class representing a single news article',
    'Source: Nested data class for the article source information',
]
for point in api_points:
    doc.add_paragraph(point, style='List Bullet')

doc.add_page_break()

add_heading_styled('NewsRepository.kt', 3)

doc.add_paragraph(
    'The Repository acts as a single source of truth for news data. It abstracts the data source '
    '(network API) from the rest of the application.'
)

repository_code = """package com.example.newsapp.data

class NewsRepository {
    private val newsApi: NewsApi

    init {
        val retrofit = Retrofit.Builder()
            .baseUrl("https://newsapi.org/v2/")
            .addConverterFactory(GsonConverterFactory.create())
            .build()

        newsApi = retrofit.create(NewsApi::class.java)
    }

    suspend fun getTopHeadlines(apiKey: String): List<Article> {
        return try {
            val response = newsApi.getTopHeadlines(apiKey = apiKey)
            if (response.status == "ok") {
                response.articles
            } else {
                emptyList()
            }
        } catch (e: Exception) {
            e.printStackTrace()
            emptyList()
        }
    }
}"""
add_code_block(repository_code)

doc.add_paragraph('Key Features:')
repo_points = [
    'Retrofit Configuration: Sets up the base URL and JSON converter',
    'Error Handling: Catches exceptions and returns empty list on failure',
    'Suspend Function: Allows calling from coroutines without blocking',
]
for point in repo_points:
    doc.add_paragraph(point, style='List Bullet')

doc.add_page_break()

# 6.3 UI Layer
add_heading_styled('6.3 UI Layer', 2)

add_heading_styled('NewsViewModel.kt', 3)

doc.add_paragraph(
    'The ViewModel manages UI state and business logic. It holds the article list and loading state '
    'as observable StateFlow streams.'
)

viewmodel_code = """package com.example.newsapp.ui

class NewsViewModel : ViewModel() {
    private val repository = NewsRepository()

    private val _articles = MutableStateFlow<List<Article>>(emptyList())
    val articles: StateFlow<List<Article>> = _articles.asStateFlow()

    private val _isLoading = MutableStateFlow(false)
    val isLoading: StateFlow<Boolean> = _isLoading.asStateFlow()

    init {
        fetchNews()
    }

    fun fetchNews() {
        viewModelScope.launch {
            _isLoading.value = true
            val apiKey = "YOUR_API_KEY"
            val result = repository.getTopHeadlines(apiKey)
            _articles.value = result
            _isLoading.value = false
        }
    }
}"""
add_code_block(viewmodel_code)

doc.add_paragraph('Key Concepts:')
vm_points = [
    'StateFlow: Provides observable state that Compose can collect',
    'MutableStateFlow: Private mutable version for internal updates',
    'viewModelScope: Coroutine scope tied to ViewModel lifecycle',
    'init block: Automatically fetches news when ViewModel is created',
]
for point in vm_points:
    doc.add_paragraph(point, style='List Bullet')

doc.add_page_break()

add_heading_styled('NewsScreen.kt', 3)

doc.add_paragraph(
    'The main screen composable that displays the list of news articles. It observes the ViewModel\'s '
    'state and renders the UI accordingly.'
)

screen_code = """@OptIn(ExperimentalMaterial3Api::class)
@Composable
fun NewsScreen(viewModel: NewsViewModel) {
    val articles by viewModel.articles.collectAsState()
    val isLoading by viewModel.isLoading.collectAsState()

    Scaffold(
        topBar = {
            TopAppBar(title = { Text("Top Headlines") })
        }
    ) { paddingValues ->
        Box(modifier = Modifier.padding(paddingValues).fillMaxSize()) {
            if (isLoading) {
                CircularProgressIndicator(
                    modifier = Modifier.align(Alignment.Center)
                )
            } else {
                LazyColumn(
                    contentPadding = PaddingValues(16.dp),
                    verticalArrangement = Arrangement.spacedBy(16.dp)
                ) {
                    items(articles) { article ->
                        ArticleItem(article)
                    }
                }
            }
        }
    }
}"""
add_code_block(screen_code)

add_heading_styled('ArticleItem Composable', 3)

article_code = """@Composable
fun ArticleItem(article: Article) {
    val context = LocalContext.current
    
    Card(
        modifier = Modifier
            .fillMaxWidth()
            .clickable {
                val intent = Intent(Intent.ACTION_VIEW, Uri.parse(article.url))
                context.startActivity(intent)
            },
        elevation = CardDefaults.cardElevation(defaultElevation = 4.dp)
    ) {
        Column {
            // Article image
            if (!article.urlToImage.isNullOrEmpty()) {
                AsyncImage(
                    model = article.urlToImage,
                    contentDescription = null,
                    modifier = Modifier.fillMaxWidth().height(200.dp),
                    contentScale = ContentScale.Crop
                )
            }
            // Article content
            Column(modifier = Modifier.padding(16.dp)) {
                Text(text = article.title, 
                     style = MaterialTheme.typography.titleMedium,
                     fontWeight = FontWeight.Bold)
                     
                if (!article.description.isNullOrEmpty()) {
                    Text(text = article.description,
                         maxLines = 3,
                         overflow = TextOverflow.Ellipsis)
                }
                
                Text(text = article.source.name,
                     color = MaterialTheme.colorScheme.primary)
            }
        }
    }
}"""
add_code_block(article_code)

doc.add_page_break()

# 6.4 Theme
add_heading_styled('6.4 Theme', 2)

doc.add_paragraph(
    'The theme configuration enables Material 3 theming with support for light mode, dark mode, '
    'and dynamic colors on Android 12+.'
)

add_heading_styled('Theme.kt', 3)

theme_code = """@Composable
fun NewsAppTheme(
    darkTheme: Boolean = isSystemInDarkTheme(),
    dynamicColor: Boolean = true,
    content: @Composable () -> Unit
) {
    val colorScheme = when {
        dynamicColor && Build.VERSION.SDK_INT >= Build.VERSION_CODES.S -> {
            val context = LocalContext.current
            if (darkTheme) dynamicDarkColorScheme(context) 
            else dynamicLightColorScheme(context)
        }
        darkTheme -> DarkColorScheme
        else -> LightColorScheme
    }
    
    val view = LocalView.current
    if (!view.isInEditMode) {
        SideEffect {
            val window = (view.context as Activity).window
            window.statusBarColor = colorScheme.primary.toArgb()
            WindowCompat.getInsetsController(window, view)
                .isAppearanceLightStatusBars = darkTheme
        }
    }

    MaterialTheme(
        colorScheme = colorScheme,
        typography = MaterialTheme.typography,
        content = content
    )
}"""
add_code_block(theme_code)

doc.add_paragraph('Theme Features:')
theme_points = [
    'Dynamic Colors: Uses wallpaper-based colors on Android 12+',
    'Dark Mode: Automatically switches based on system settings',
    'Status Bar: Updates status bar color to match theme',
    'Fallback: Uses predefined color schemes on older Android versions',
]
for point in theme_points:
    doc.add_paragraph(point, style='List Bullet')

doc.add_page_break()

# ============================================================
# 7. API INTEGRATION
# ============================================================
add_heading_styled('7. API Integration', 1)

doc.add_paragraph(
    'The application integrates with NewsAPI.org to fetch real-time news headlines. NewsAPI provides '
    'access to news articles from thousands of sources worldwide.'
)

add_heading_styled('API Configuration', 2)

api_config_table = doc.add_table(rows=4, cols=2)
api_config_table.style = 'Table Grid'
api_config = [
    ('Base URL', 'https://newsapi.org/v2/'),
    ('Endpoint', '/top-headlines'),
    ('Method', 'GET'),
    ('Response Format', 'JSON'),
]
for i, (key, value) in enumerate(api_config):
    api_config_table.rows[i].cells[0].text = key
    api_config_table.rows[i].cells[1].text = value
    set_cell_shading(api_config_table.rows[i].cells[0], 'F1F5F9')

doc.add_paragraph()

add_heading_styled('Request Parameters', 2)

params_table = doc.add_table(rows=3, cols=3)
params_table.style = 'Table Grid'
params = [
    ('Parameter', 'Required', 'Description'),
    ('country', 'Yes', 'Two-letter country code (e.g., "us" for United States)'),
    ('apiKey', 'Yes', 'Your API authentication key from NewsAPI.org'),
]
for i, (param, req, desc) in enumerate(params):
    params_table.rows[i].cells[0].text = param
    params_table.rows[i].cells[1].text = req
    params_table.rows[i].cells[2].text = desc
    if i == 0:
        for cell in params_table.rows[i].cells:
            cell.paragraphs[0].runs[0].bold = True
            set_cell_shading(cell, '4F46E5')
            cell.paragraphs[0].runs[0].font.color.rgb = RGBColor(255, 255, 255)

doc.add_paragraph()

add_heading_styled('API Response Structure', 2)

response_code = """{
    "status": "ok",
    "totalResults": 38,
    "articles": [
        {
            "source": {
                "id": "bbc-news",
                "name": "BBC News"
            },
            "author": "John Doe",
            "title": "Article Title Here",
            "description": "Brief description of the article...",
            "url": "https://www.bbc.com/news/...",
            "urlToImage": "https://ichef.bbci.co.uk/...",
            "publishedAt": "2024-01-15T10:30:00Z",
            "content": "Full article content..."
        }
    ]
}"""
add_code_block(response_code, "json")

add_heading_styled('Free Tier Limitations', 2)
limitations = [
    'Maximum 100 requests per day',
    'Only works during development (not for production apps)',
    'Headlines may be delayed by a few hours',
    'Limited to certain sources',
]
for limit in limitations:
    doc.add_paragraph(limit, style='List Bullet')

doc.add_page_break()

# ============================================================
# 8. USER INTERFACE
# ============================================================
add_heading_styled('8. User Interface', 1)

doc.add_paragraph(
    'The user interface is built entirely with Jetpack Compose, providing a modern, reactive UI '
    'that automatically updates when the underlying data changes.'
)

add_heading_styled('UI Components', 2)

ui_components = [
    ('TopAppBar', 'Displays the app title "Top Headlines" at the top of the screen'),
    ('LazyColumn', 'Efficiently renders a scrollable list of article cards'),
    ('Card', 'Material 3 card component for each news article'),
    ('AsyncImage', 'Coil component for loading article images from URLs'),
    ('CircularProgressIndicator', 'Shows loading state while fetching news'),
]

for component, desc in ui_components:
    p = doc.add_paragraph()
    p.add_run(f'{component}: ').bold = True
    p.add_run(desc)

add_heading_styled('Article Card Layout', 2)

doc.add_paragraph('Each article card displays the following information:')

card_elements = [
    ('Image', 'Article thumbnail (200dp height, cropped to fit)'),
    ('Title', 'Article headline in bold, medium typography'),
    ('Description', 'Brief summary, limited to 3 lines with ellipsis'),
    ('Source', 'News source name displayed in primary theme color'),
]

card_table = doc.add_table(rows=len(card_elements)+1, cols=2)
card_table.style = 'Table Grid'
card_table.rows[0].cells[0].text = 'Element'
card_table.rows[0].cells[1].text = 'Description'
for cell in card_table.rows[0].cells:
    cell.paragraphs[0].runs[0].bold = True
    set_cell_shading(cell, '4F46E5')
    cell.paragraphs[0].runs[0].font.color.rgb = RGBColor(255, 255, 255)

for i, (element, desc) in enumerate(card_elements):
    card_table.rows[i+1].cells[0].text = element
    card_table.rows[i+1].cells[1].text = desc

add_heading_styled('User Interactions', 2)

interactions = [
    'Scrolling: Users can scroll through the list of articles using standard swipe gestures',
    'Tapping: Tapping on an article card opens the full article in the device\'s default web browser',
    'Automatic Refresh: News is fetched automatically when the app starts',
]
for interaction in interactions:
    doc.add_paragraph(interaction, style='List Bullet')

doc.add_page_break()

# ============================================================
# 9. CONFIGURATION & SETUP
# ============================================================
add_heading_styled('9. Configuration & Setup', 1)

add_heading_styled('Prerequisites', 2)

prereqs = [
    'Android Studio (latest version recommended)',
    'Android SDK 34 or higher',
    'Kotlin 1.9.x',
    'A NewsAPI.org account and API key',
]
for prereq in prereqs:
    doc.add_paragraph(prereq, style='List Bullet')

add_heading_styled('Setup Instructions', 2)

setup_steps = [
    ('Clone or Open Project', 'Open the project folder in Android Studio'),
    ('Get API Key', 'Register at newsapi.org and obtain a free API key'),
    ('Configure API Key', 'Open NewsViewModel.kt and replace the apiKey value with your key'),
    ('Sync Project', 'Click "Sync Project with Gradle Files" in Android Studio'),
    ('Run Application', 'Connect a device or start an emulator, then click Run'),
]

for i, (step, desc) in enumerate(setup_steps, 1):
    p = doc.add_paragraph()
    p.add_run(f'{i}. {step}: ').bold = True
    p.add_run(desc)

add_heading_styled('AndroidManifest.xml', 2)

doc.add_paragraph('The app requires the following permission:')

manifest_code = """<manifest xmlns:android="http://schemas.android.com/apk/res/android">
    <uses-permission android:name="android.permission.INTERNET" />
    
    <application
        android:icon="@mipmap/ic_launcher"
        android:label="@string/app_name"
        android:theme="@style/Theme.NewsApp">
        <activity
            android:name=".MainActivity"
            android:exported="true">
            <intent-filter>
                <action android:name="android.intent.action.MAIN" />
                <category android:name="android.intent.category.LAUNCHER" />
            </intent-filter>
        </activity>
    </application>
</manifest>"""
add_code_block(manifest_code, "xml")

doc.add_page_break()

# ============================================================
# 10. FUTURE IMPROVEMENTS
# ============================================================
add_heading_styled('10. Future Improvements', 1)

doc.add_paragraph(
    'The following features are planned or suggested for future development to enhance the application:'
)

improvements = [
    ('Search Functionality', 'Allow users to search for news on specific topics or keywords'),
    ('Category Filtering', 'Filter news by categories such as Sports, Technology, Business, Entertainment'),
    ('Favorites/Bookmarks', 'Save articles to read later, stored locally'),
    ('Offline Mode', 'Cache downloaded articles for reading without internet connection'),
    ('Pull to Refresh', 'Swipe down gesture to refresh the news feed'),
    ('Share Feature', 'Share interesting articles via social media or messaging apps'),
    ('Multiple Countries', 'Allow users to select news from different countries'),
    ('Push Notifications', 'Get alerts for breaking news stories'),
    ('Article Reading Mode', 'In-app article viewer instead of opening external browser'),
    ('Pagination', 'Load more articles as user scrolls (infinite scrolling)'),
]

for i, (feature, desc) in enumerate(improvements, 1):
    p = doc.add_paragraph()
    p.add_run(f'{i}. {feature}: ').bold = True
    p.add_run(desc)

doc.add_page_break()

# ============================================================
# 11. APPENDIX
# ============================================================
add_heading_styled('11. Appendix', 1)

add_heading_styled('A. Build Configuration', 2)

doc.add_paragraph('Key build.gradle.kts settings:')

build_config = """android {
    namespace = "com.example.newsapp"
    compileSdk = 34

    defaultConfig {
        applicationId = "com.example.newsapp"
        minSdk = 24
        targetSdk = 34
        versionCode = 1
        versionName = "1.0"
    }

    buildFeatures {
        compose = true
    }
    
    composeOptions {
        kotlinCompilerExtensionVersion = "1.5.1"
    }
}"""
add_code_block(build_config)

add_heading_styled('B. Color Definitions', 2)

color_code = """// Color.kt
val Purple80 = Color(0xFFD0BCFF)
val PurpleGrey80 = Color(0xFFCCC2DC)
val Pink80 = Color(0xFFEFB8C8)

val Purple40 = Color(0xFF6650a4)
val PurpleGrey40 = Color(0xFF625b71)
val Pink40 = Color(0xFF7D5260)"""
add_code_block(color_code)

add_heading_styled('C. Useful Resources', 2)

resources = [
    ('Jetpack Compose Documentation', 'https://developer.android.com/jetpack/compose'),
    ('Material 3 Design', 'https://m3.material.io/'),
    ('NewsAPI Documentation', 'https://newsapi.org/docs'),
    ('Kotlin Documentation', 'https://kotlinlang.org/docs/'),
    ('Retrofit Documentation', 'https://square.github.io/retrofit/'),
]

for name, url in resources:
    p = doc.add_paragraph()
    p.add_run(f'• {name}: ').bold = True
    p.add_run(url)

# ============================================================
# SAVE DOCUMENT
# ============================================================
output_path = r"c:\Users\LenovoPC\Desktop\first_app\NewsApp_Full_Documentation.docx"
doc.save(output_path)

print(f"\n✅ Documentation saved successfully!")
print(f"📁 Location: {output_path}")
print(f"\n📄 Document contains the following sections:")
print("   1. Introduction")
print("   2. Project Overview")
print("   3. Technologies Used")
print("   4. Project Structure")
print("   5. Architecture")
print("   6. Core Components")
print("   7. API Integration")
print("   8. User Interface")
print("   9. Configuration & Setup")
print("   10. Future Improvements")
print("   11. Appendix")
